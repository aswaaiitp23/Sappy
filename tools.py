"""
Tool definitions and dispatch.

Each tool has two parts:
  1. A JSON schema in TOOLS — sent to Anthropic so the model knows what's available.
  2. A Python implementation below — called by agent.py's loop after the model
     requests the tool. The model never runs Python; your loop does.

dispatch_tool(name, inputs) -> str  is the single entry point used by agent.py.
"""

import csv
import os
from pathlib import Path

import openpyxl
import pdfplumber

# Maximum characters returned by read_file before truncation.
# Keeps individual tool results from blowing out the context window.
DEFAULT_MAX_CHARS = 20_000

# ── tool schemas (sent to Anthropic) ─────────────────────────────────────────

TOOLS = [
    {
        "name": "read_file",
        "description": (
            "Read a file and return its text content. "
            "Supports PDF, Excel (.xlsx), CSV, and plain text (.txt, .md). "
            "For spreadsheets, returns all sheets as tab-separated text. "
            "For PDFs, returns extracted text page by page. "
            "Large files are truncated to max_chars characters."
        ),
        "input_schema": {
            "type": "object",
            "properties": {
                "path": {
                    "type": "string",
                    "description": "Path to the file.",
                },
                "max_chars": {
                    "type": "integer",
                    "description": (
                        f"Truncate output to this many characters "
                        f"(default {DEFAULT_MAX_CHARS}). Increase if you need more."
                    ),
                },
            },
            "required": ["path"],
        },
    },
    {
        "name": "list_files",
        "description": "List files in a directory, showing names and sizes.",
        "input_schema": {
            "type": "object",
            "properties": {
                "directory": {
                    "type": "string",
                    "description": "Path to the directory to list.",
                }
            },
            "required": ["directory"],
        },
    },
    {
        "name": "load_skill",
        "description": (
            "Load the full content of a named skill from the ./skills folder. "
            "Use this when you need detailed step-by-step instructions for a "
            "specific task type (e.g. financial_analysis)."
        ),
        "input_schema": {
            "type": "object",
            "properties": {
                "name": {
                    "type": "string",
                    "description": "Skill filename without the .md extension.",
                }
            },
            "required": ["name"],
        },
    },
]

# ── tool implementations ──────────────────────────────────────────────────────

def read_file(path: str, max_chars: int = DEFAULT_MAX_CHARS) -> str:
    if not os.path.exists(path):
        return f"Error: file not found: {path}"

    ext = os.path.splitext(path)[1].lower()
    try:
        if ext == ".pdf":
            raw = _read_pdf(path)
        elif ext in (".xlsx", ".xls"):
            raw = _read_excel(path)
        elif ext == ".csv":
            raw = _read_csv(path)
        elif ext in (".txt", ".md", ".json", ".yaml", ".yml", ".toml", ""):
            raw = Path(path).read_text(encoding="utf-8", errors="replace")
        else:
            return f"Unsupported file type '{ext}'. Supported: pdf, xlsx, csv, txt, md."
    except Exception as exc:
        return f"Error reading {path}: {exc}"

    if len(raw) > max_chars:
        return raw[:max_chars] + f"\n\n… [truncated — {len(raw) - max_chars} chars omitted]"
    return raw


def _read_pdf(path: str) -> str:
    pages = []
    with pdfplumber.open(path) as pdf:
        for i, page in enumerate(pdf.pages):
            text = page.extract_text()
            if text:
                pages.append(f"--- Page {i + 1} ---\n{text}")
    return "\n\n".join(pages) if pages else "No extractable text found in PDF."


def _read_excel(path: str) -> str:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    parts = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows = []
        for row in ws.iter_rows(values_only=True):
            if any(cell is not None for cell in row):
                rows.append("\t".join("" if c is None else str(c) for c in row))
        parts.append(f"=== Sheet: {sheet_name} ===\n" + "\n".join(rows))
    return "\n\n".join(parts)


def _read_csv(path: str) -> str:
    with open(path, newline="", encoding="utf-8", errors="replace") as f:
        rows = ["\t".join(row) for row in csv.reader(f)]
    return "\n".join(rows)


def list_files(directory: str) -> str:
    if not os.path.exists(directory):
        return f"Error: directory not found: {directory}"
    if not os.path.isdir(directory):
        return f"Error: not a directory: {directory}"

    entries = []
    for name in sorted(os.listdir(directory)):
        full = os.path.join(directory, name)
        size = os.path.getsize(full)
        kind = "dir" if os.path.isdir(full) else f"{size:,} bytes"
        entries.append(f"  {name:<45} {kind}")

    if not entries:
        return f"Empty directory: {directory}"
    return f"Contents of {directory}:\n" + "\n".join(entries)


def load_skill(name: str) -> str:
    skill_path = os.path.join("skills", f"{name}.md")
    if not os.path.exists(skill_path):
        available = [
            f[:-3] for f in os.listdir("skills") if f.endswith(".md")
        ] if os.path.isdir("skills") else []
        return f"Skill '{name}' not found. Available skills: {available}"
    with open(skill_path, encoding="utf-8") as f:
        return f.read()


# ── dispatch table ────────────────────────────────────────────────────────────

def dispatch_tool(name: str, inputs: dict) -> str:
    """Route a tool call by name. Called by the agent loop — never by the model."""
    if name == "read_file":
        return read_file(inputs["path"], inputs.get("max_chars", DEFAULT_MAX_CHARS))
    if name == "list_files":
        return list_files(inputs["directory"])
    if name == "load_skill":
        return load_skill(inputs["name"])
    return f"Unknown tool: '{name}'. Available: read_file, list_files, load_skill."
